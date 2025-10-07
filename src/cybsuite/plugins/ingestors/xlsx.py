from pathlib import Path

from cybsuite.cyberdb import BaseIngestor, Metadata
from openpyxl import load_workbook


class XlsxIngestor(BaseIngestor):
    name = "xlsx"

    metadata = Metadata(
        description="Ingest Excel files (.xlsx) with different sheets feeding different entity types based on sheet names."
    )

    autodetect_is_file = True

    @classmethod
    def autodetect_from_path(cls, path: Path) -> bool:
        return path.suffix.lower() == ".cybsuite.xlsx"

    def do_run(self, filepath):
        workbook = load_workbook(filepath, read_only=True)

        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]

            # Skip empty sheets
            if worksheet.max_row <= 1:
                continue

            # Get headers from first row
            headers = []
            for cell in worksheet[1]:
                headers.append(cell.value if cell.value else "")

            # Process each row starting from row 2
            for row in worksheet.iter_rows(min_row=2, values_only=True):
                if not any(row):  # Skip empty rows
                    continue

                # Create a dict from headers and row values
                row_data = dict(zip(headers, row))

                # Remove None values
                row_data = {k: v for k, v in row_data.items() if v is not None}

                # Feed using sheet name as entity type
                # TODO: cast types (int, float, bool, date, time, datetime, dict, list, set)
                self.cyberdb.feed(sheet_name.lower(), **row_data)

        workbook.close()
